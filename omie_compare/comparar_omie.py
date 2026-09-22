"""Compara Categorias e Departamentos entre vários apps (empresas) da OMIE
e verifica se cada código tem lançamento (Contas a Pagar / Contas a Receber)
dentro de um período configurável.

Gera `omie_comparacao_<data>.xlsx` e salva o JSON bruto de cada chamada em
`raw/<app>_<tipo>.json` para auditoria manual da heurística de detecção de
campos (veja o README.md).
"""

from __future__ import annotations

import json
import os
from datetime import datetime, timedelta
from typing import Any

import yaml
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from nadarte_plano_de_contas.omie_compare.omie_client import OmieClient, OmieError, find_records_list, normalize_key

# ---------------------------------------------------------------------------
# CONFIGURAÇÃO DO PERÍODO
# ---------------------------------------------------------------------------
# Ajuste PERIODO_DIAS (ou defina DATA_DE_STR/DATA_ATE_STR manualmente) para
# controlar o intervalo de Contas a Pagar/Receber consultado.
PERIODO_DIAS = 365

_hoje = datetime.today()
DATA_DE_STR = (_hoje - timedelta(days=PERIODO_DIAS)).strftime("%d/%m/%Y")
DATA_ATE_STR = _hoje.strftime("%d/%m/%Y")

CONFIG_PATH = "config.yaml"
RAW_DIR = "raw"
REGISTROS_POR_PAGINA_CADASTRO = 200
REGISTROS_POR_PAGINA_CONTAS = 200

YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
BOLD_FONT = Font(bold=True)


# ---------------------------------------------------------------------------
# Helpers de detecção dinâmica de campos
# ---------------------------------------------------------------------------

def find_code_field(record: dict[str, Any]) -> str | None:
    """Acha a chave de 'código' de um registro de categoria/departamento,
    evitando pegar 'codigo_pai' ou outros códigos secundários."""
    norm_map = {k: normalize_key(k) for k in record.keys()}
    for k, n in norm_map.items():
        if n == "codigo":
            return k
    for k, n in norm_map.items():
        if "codigo" in n and "pai" not in n:
            return k
    for k, n in norm_map.items():
        if n.startswith("cod") and "pai" not in n:
            return k
    return None


def find_description_field(record: dict[str, Any]) -> str | None:
    """Acha a chave de 'descrição' de um registro de categoria/departamento."""
    norm_map = {k: normalize_key(k) for k in record.keys()}
    for preferred in ("descricao_padrao", "descricao"):
        for k, n in norm_map.items():
            if n == preferred:
                return k
    for k, n in norm_map.items():
        if "descricao" in n:
            return k
    for k, n in norm_map.items():
        if "nome" in n:
            return k
    return None


def collect_used_codes(obj: Any, keyword_substr: str) -> set[str]:
    """Percorre recursivamente `obj` (um título financeiro) coletando os
    valores de qualquer chave cujo nome normalizado contenha
    `keyword_substr` (ex.: "categ" ou "departamento"), sem assumir um
    caminho fixo como distribuicao[].cCodCateg."""
    codes: set[str] = set()

    def _walk(node: Any) -> None:
        if isinstance(node, dict):
            for k, v in node.items():
                if keyword_substr in normalize_key(k):
                    if isinstance(v, (str, int, float)) and str(v).strip():
                        codes.add(str(v).strip())
                    elif isinstance(v, list):
                        for item in v:
                            if isinstance(item, (str, int, float)) and str(item).strip():
                                codes.add(str(item).strip())
                _walk(v)
        elif isinstance(node, list):
            for item in node:
                _walk(item)

    _walk(obj)
    return codes


# ---------------------------------------------------------------------------
# Chamadas à API
# ---------------------------------------------------------------------------

def fetch_all_pages(
    client: OmieClient,
    resource: str,
    method: str,
    base_params: dict[str, Any],
    per_page: int,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Percorre todas as páginas de um endpoint, retornando
    (registros_mesclados, paginas_brutas)."""
    raw_pages: list[dict[str, Any]] = []
    records: list[dict[str, Any]] = []
    pagina = 1
    total_paginas = 1
    while pagina <= total_paginas:
        params = dict(base_params)
        params["pagina"] = pagina
        params["registros_por_pagina"] = per_page
        data = client.call(resource, method, params)
        raw_pages.append(data)
        records.extend(find_records_list(data))
        total_paginas = int(data.get("total_de_paginas") or data.get("totalDePaginas") or 1)
        pagina += 1
    return records, raw_pages


def save_raw(app_name: str, kind: str, raw_pages: list[dict[str, Any]]) -> None:
    os.makedirs(RAW_DIR, exist_ok=True)
    path = os.path.join(RAW_DIR, f"{app_name}_{kind}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(raw_pages, f, ensure_ascii=False, indent=2)


def build_cadastro_dict(records: list[dict[str, Any]]) -> dict[str, str]:
    result: dict[str, str] = {}
    for rec in records:
        code_field = find_code_field(rec)
        if code_field is None:
            continue
        code = str(rec[code_field]).strip()
        if not code:
            continue
        desc_field = find_description_field(rec)
        desc = str(rec[desc_field]).strip() if desc_field else ""
        result[code] = desc
    return result


# ---------------------------------------------------------------------------
# Geração da planilha
# ---------------------------------------------------------------------------

def autofit_columns(ws) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, 10), 60)


def bold_header(ws) -> None:
    for cell in ws[1]:
        cell.font = BOLD_FONT


def write_cadastro_sheet(writer, sheet_name: str, cadastros: dict[str, dict[str, str]], app_names: list[str]) -> None:
    import pandas as pd

    all_codes = sorted({code for cad in cadastros.values() for code in cad.keys()})
    rows = []
    divergent_flags = []
    for code in all_codes:
        row = {"codigo": code}
        descricoes_presentes = set()
        apps_com_codigo = 0
        for app in app_names:
            desc = cadastros.get(app, {}).get(code)
            row[app] = desc if desc is not None else ""
            if desc is not None:
                apps_com_codigo += 1
                descricoes_presentes.add(desc)
        rows.append(row)
        diverge = len(descricoes_presentes) > 1 or apps_com_codigo < len(app_names)
        divergent_flags.append(diverge)

    df = pd.DataFrame(rows, columns=["codigo"] + app_names)
    df.to_excel(writer, sheet_name=sheet_name, index=False)

    ws = writer.sheets[sheet_name]
    bold_header(ws)
    for row_idx, diverge in enumerate(divergent_flags, start=2):
        if diverge:
            for col_idx in range(1, len(app_names) + 2):
                ws.cell(row=row_idx, column=col_idx).fill = YELLOW_FILL
    autofit_columns(ws)


def write_uso_sheet(
    writer,
    sheet_name: str,
    cadastros: dict[str, dict[str, str]],
    usados: dict[str, set[str]],
    app_names: list[str],
) -> None:
    import pandas as pd

    all_codes = sorted({code for cad in cadastros.values() for code in cad.keys()})
    rows = []
    for code in all_codes:
        row = {"codigo": code}
        for app in app_names:
            if code not in cadastros.get(app, {}):
                status = "Não existe"
            elif code in usados.get(app, set()):
                status = "Com lançamento"
            else:
                status = "Sem lançamento"
            row[app] = status
        rows.append(row)

    df = pd.DataFrame(rows, columns=["codigo"] + app_names)
    df.to_excel(writer, sheet_name=sheet_name, index=False)

    ws = writer.sheets[sheet_name]
    bold_header(ws)
    autofit_columns(ws)


def write_resumo_sheet(
    writer,
    app_names: list[str],
    categorias_cadastro: dict[str, dict[str, str]],
    departamentos_cadastro: dict[str, dict[str, str]],
    categorias_usadas: dict[str, set[str]],
    departamentos_usadas: dict[str, set[str]],
    erros: dict[str, str],
) -> None:
    import pandas as pd

    rows = []
    for app in app_names:
        rows.append(
            {
                "app": app,
                "categorias_cadastradas": len(categorias_cadastro.get(app, {})),
                "categorias_com_lancamento": len(
                    set(categorias_cadastro.get(app, {})) & categorias_usadas.get(app, set())
                ),
                "departamentos_cadastrados": len(departamentos_cadastro.get(app, {})),
                "departamentos_com_lancamento": len(
                    set(departamentos_cadastro.get(app, {})) & departamentos_usadas.get(app, set())
                ),
                "erro": erros.get(app, ""),
            }
        )
    df = pd.DataFrame(rows)
    df.to_excel(writer, sheet_name="Resumo", index=False)
    ws = writer.sheets["Resumo"]
    bold_header(ws)
    autofit_columns(ws)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        config = yaml.safe_load(f)

    apps_config = config["apps"]
    app_names = list(apps_config.keys())

    categorias_cadastro: dict[str, dict[str, str]] = {}
    departamentos_cadastro: dict[str, dict[str, str]] = {}
    categorias_usadas: dict[str, set[str]] = {}
    departamentos_usadas: dict[str, set[str]] = {}
    erros: dict[str, str] = {}

    for app_name in app_names:
        creds = apps_config[app_name]
        print(f"\n=== {app_name} ===")
        client = OmieClient(creds["app_key"], creds["app_secret"])

        try:
            print("  Buscando categorias...")
            cat_records, cat_raw = fetch_all_pages(
                client, "geral/categorias", "ListarCategorias", {}, REGISTROS_POR_PAGINA_CADASTRO
            )
            save_raw(app_name, "categorias", cat_raw)
            categorias_cadastro[app_name] = build_cadastro_dict(cat_records)
            print(f"    {len(categorias_cadastro[app_name])} categorias cadastradas")

            print("  Buscando departamentos...")
            dep_records, dep_raw = fetch_all_pages(
                client, "geral/departamentos", "ListarDepartamentos", {}, REGISTROS_POR_PAGINA_CADASTRO
            )
            save_raw(app_name, "departamentos", dep_raw)
            departamentos_cadastro[app_name] = build_cadastro_dict(dep_records)
            print(f"    {len(departamentos_cadastro[app_name])} departamentos cadastrados")

            periodo_params = {
                "filtrar_por_data_de": DATA_DE_STR,
                "filtrar_por_data_ate": DATA_ATE_STR,
            }

            print(f"  Buscando contas a pagar ({DATA_DE_STR} a {DATA_ATE_STR})...")
            pagar_records, pagar_raw = fetch_all_pages(
                client, "financas/contapagar", "ListarContasPagar", periodo_params, REGISTROS_POR_PAGINA_CONTAS
            )
            save_raw(app_name, "contas_pagar", pagar_raw)
            print(f"    {len(pagar_records)} títulos a pagar")

            print(f"  Buscando contas a receber ({DATA_DE_STR} a {DATA_ATE_STR})...")
            receber_records, receber_raw = fetch_all_pages(
                client, "financas/contareceber", "ListarContasReceber", periodo_params, REGISTROS_POR_PAGINA_CONTAS
            )
            save_raw(app_name, "contas_receber", receber_raw)
            print(f"    {len(receber_records)} títulos a receber")

            usados_categ: set[str] = set()
            usados_dep: set[str] = set()
            for titulo in pagar_records + receber_records:
                usados_categ |= collect_used_codes(titulo, "categ")
                usados_dep |= collect_used_codes(titulo, "departamento")
            categorias_usadas[app_name] = usados_categ
            departamentos_usadas[app_name] = usados_dep
            print(
                f"    {len(usados_categ)} categorias com lançamento, "
                f"{len(usados_dep)} departamentos com lançamento"
            )

        except OmieError as exc:
            msg = f"{exc.faultcode}: {exc.faultstring}"
            print(f"  ERRO OMIE: {msg}")
            erros[app_name] = msg
            categorias_cadastro.setdefault(app_name, {})
            departamentos_cadastro.setdefault(app_name, {})
            categorias_usadas.setdefault(app_name, set())
            departamentos_usadas.setdefault(app_name, set())
        except Exception as exc:  # noqa: BLE001
            msg = str(exc)
            print(f"  ERRO: {msg}")
            erros[app_name] = msg
            categorias_cadastro.setdefault(app_name, {})
            departamentos_cadastro.setdefault(app_name, {})
            categorias_usadas.setdefault(app_name, set())
            departamentos_usadas.setdefault(app_name, set())

    import pandas as pd  # local import to keep module import fast/optional for --help-like usage

    output_path = f"omie_comparacao_{datetime.today().strftime('%Y%m%d_%H%M%S')}.xlsx"
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        write_resumo_sheet(
            writer, app_names, categorias_cadastro, departamentos_cadastro,
            categorias_usadas, departamentos_usadas, erros,
        )
        write_cadastro_sheet(writer, "Categorias", categorias_cadastro, app_names)
        write_cadastro_sheet(writer, "Departamentos", departamentos_cadastro, app_names)
        write_uso_sheet(writer, "Uso_Categorias", categorias_cadastro, categorias_usadas, app_names)
        write_uso_sheet(writer, "Uso_Departamentos", departamentos_cadastro, departamentos_usadas, app_names)

    print(f"\nPlanilha gerada: {output_path}")
    if erros:
        print("\nApps com erro:")
        for app, msg in erros.items():
            print(f"  {app}: {msg}")


if __name__ == "__main__":
    main()
